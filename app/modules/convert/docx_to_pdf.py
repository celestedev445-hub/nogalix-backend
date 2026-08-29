from __future__ import annotations

import io
from typing import Any, Iterator
from xml.sax.saxutils import escape

import pymupdf
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run


WORD_MARGIN_PT = 72  # 2,54 cm — même marge Word « Normal »


def _align_css(alignment: WD_ALIGN_PARAGRAPH | None) -> str:
    if alignment == WD_ALIGN_PARAGRAPH.CENTER:
        return "center"
    if alignment == WD_ALIGN_PARAGRAPH.RIGHT:
        return "right"
    if alignment == WD_ALIGN_PARAGRAPH.JUSTIFY:
        return "justify"
    return "left"


def _run_color(run: Run) -> str | None:
    try:
        rgb = run.font.color.rgb
    except Exception:
        return None
    if rgb is None:
        return None
    return f"#{str(rgb)}"


def _run_html(run: Run) -> str:
    text = escape(run.text or "").replace("\n", "<br/>")
    if not text:
        return ""
    styles: list[str] = []
    if run.bold:
        styles.append("font-weight:700")
    if run.italic:
        styles.append("font-style:italic")
    if run.underline:
        styles.append("text-decoration:underline")
    size = run.font.size
    if size is not None:
        styles.append(f"font-size:{float(size.pt):.1f}pt")
    name = run.font.name
    if name:
        styles.append(f"font-family:'{escape(name)}',sans-serif")
    color = _run_color(run)
    if color:
        styles.append(f"color:{color}")
    style_attr = f' style="{";".join(styles)}"' if styles else ""
    return f"<span{style_attr}>{text}</span>"


def _paragraph_html(paragraph: Paragraph) -> str:
    inner = "".join(_run_html(run) for run in paragraph.runs)
    if not inner.strip():
        return "<p>&nbsp;</p>"
    align = _align_css(paragraph.alignment)
    style = f"text-align:{align};margin:0 0 8pt 0;line-height:1.15"
    style_name = (paragraph.style.name or "") if paragraph.style is not None else ""
    if style_name.startswith("Heading 1") or style_name == "Title":
        return f'<h1 style="{style};font-size:20pt">{inner}</h1>'
    if style_name.startswith("Heading 2"):
        return f'<h2 style="{style};font-size:16pt">{inner}</h2>'
    if style_name.startswith("Heading 3"):
        return f'<h3 style="{style};font-size:14pt">{inner}</h3>'
    return f'<p style="{style}">{inner}</p>'


def _cell_html(cell: _Cell) -> str:
    parts = [_paragraph_html(paragraph) for paragraph in cell.paragraphs]
    return (
        "<td style='border:0.5pt solid #cbd5e1;padding:4pt 6pt;vertical-align:top'>"
        + "".join(parts)
        + "</td>"
    )


def _table_html(table: Table) -> str:
    rows: list[str] = []
    for row in table.rows:
        cells = "".join(_cell_html(cell) for cell in row.cells)
        rows.append(f"<tr>{cells}</tr>")
    return (
        "<table style='width:100%;border-collapse:collapse;margin:0 0 12pt 0'>"
        + "".join(rows)
        + "</table>"
    )


def _iter_block_items(document: Document) -> Iterator[tuple[str, Any]]:
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            yield "p", child
        elif child.tag == qn("w:tbl"):
            yield "tbl", child


def _docx_to_html(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    paragraph_map = {p._element: p for p in document.paragraphs}
    table_map = {t._element: t for t in document.tables}
    chunks: list[str] = [
        '<html><body style="font-family:Calibri,Arial,sans-serif;font-size:11pt;color:#0F172A">'
    ]
    for kind, element in _iter_block_items(document):
        if kind == "p":
            paragraph = paragraph_map.get(element)
            if paragraph is not None:
                chunks.append(_paragraph_html(paragraph))
        else:
            table = table_map.get(element)
            if table is not None:
                chunks.append(_table_html(table))
    chunks.append("</body></html>")
    return "".join(chunks)


def convert_docx_to_pdf(content: bytes) -> bytes:
    if not content:
        raise ValueError("Fichier Word vide.")

    try:
        html_doc = _docx_to_html(content)
    except Exception as exc:
        raise ValueError("Impossible de lire ce fichier Word.") from exc

    if "<p" not in html_doc and "<h" not in html_doc and "<table" not in html_doc:
        raise ValueError("Ce document Word ne contient pas de texte à convertir.")

    mediabox = pymupdf.paper_rect("a4")
    where = pymupdf.Rect(
        WORD_MARGIN_PT,
        WORD_MARGIN_PT,
        mediabox.width - WORD_MARGIN_PT,
        mediabox.height - WORD_MARGIN_PT,
    )
    story = pymupdf.Story(html_doc)
    buffer = io.BytesIO()
    writer = pymupdf.DocumentWriter(buffer)
    more = True
    pages = 0
    while more:
        device = writer.begin_page(mediabox)
        more, _filled = story.place(where)
        story.draw(device)
        writer.end_page()
        pages += 1
        if pages > 80:
            break
    writer.close()
    result = buffer.getvalue()
    if not result:
        raise ValueError("La conversion Word vers PDF a échoué.")
    return result
